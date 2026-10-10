import io
import streamlit as st
import google.generativeai as genai
from docx import Document

# --- CẤU HÌNH TRANG WEB ---
st.set_page_config(
    page_title="Trợ Lý AI CV 7991 & CV 5512",
    page_icon="📚",
    layout="wide",
)

st.title("📚 Trợ Lý AI Xuất Bản Kế Hoạch Bài Dạy & Bài Kiểm Tra")
st.caption("Đúng chuẩn cấu trúc Công văn 7991/BGDĐT-GDTrH & Công văn 5512/BGDĐT-GDTrH")

# --- LẤY API KEY ẨN TỪ SECRETS CỦA STREAMLIT CLOUD ---
try:
    GOOGLE_API_KEY = st.secrets["GOOGLE_API_KEY"]
except Exception:
    GOOGLE_API_KEY = None

# --- GIAO DIỆN NHẬP LIỆU CHÍNH ---
with st.form("main_form"):
    col1, col2 = st.columns(2)
    
    with col1:
        subject = st.selectbox(
            "📖 Môn học",
            ["Toán", "Ngữ văn", "Tiếng Anh", "Khoa học tự nhiên", "Vật lý", "Hóa học", "Sinh học", "Lịch sử và Địa lí", "Tin học", "Công nghệ"]
        )
        grade = st.selectbox(
            "🏫 Khối lớp",
            ["Lớp 6", "Lớp 7", "Lớp 8", "Lớp 9", "Lớp 10", "Lớp 11", "Lớp 12"]
        )
    
    with col2:
        book = st.selectbox(
            "📚 Bộ sách giáo khoa",
            ["Kết nối tri thức với cuộc sống", "Chân trời sáng tạo", "Cánh diều"]
        )
        duration = st.number_input("⏱️ Thời lượng (Tiết)", min_value=1, max_value=6, value=2)

    lesson_title = st.text_input("🎯 Tên Bài học / Chủ đề kiểm tra", placeholder="Ví dụ: Định lý Pythagore / Ôn tập chương I...")

    export_type = st.multiselect(
        "🛠️ Chọn sản phẩm cần xuất bản:",
        [
            "Đề kiểm tra định kỳ + Ma trận & Bản đặc tả (Chuẩn CV 7991)",
            "Kế hoạch bài dạy / Giáo án (Chuẩn CV 5512)",
            "Kịch bản Slide bài giảng (PowerPoint)"
        ],
        default=["Đề kiểm tra định kỳ + Ma trận & Bản đặc tả (Chuẩn CV 7991)"]
    )

    submit_btn = st.form_submit_button("🚀 Bắt đầu tạo sản phẩm AI", type="primary")

# --- HÀM TẠO FILE WORD (.DOCX) ---
def export_to_word(text_content):
    doc = Document()
    doc.add_heading("TÀI LIỆU GIÁO DỤC - CHUẨN CÔNG VẤN BGDĐT", level=0)
    
    for paragraph in text_content.split("\n"):
        p_clean = paragraph.strip()
        if not p_clean:
            continue
        if p_clean.startswith("# "):
            doc.add_heading(p_clean.replace("# ", ""), level=1)
        elif p_clean.startswith("## "):
            doc.add_heading(p_clean.replace("## ", ""), level=2)
        elif p_clean.startswith("### "):
            doc.add_heading(p_clean.replace("### ", ""), level=3)
        elif p_clean.startswith("- ") or p_clean.startswith("* "):
            doc.add_paragraph(p_clean[2:], style='List Bullet')
        else:
            doc.add_paragraph(p_clean)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# --- XỬ LÝ KHI BẤM NÚT GỌI AI THẬT ---
if submit_btn:
    if not GOOGLE_API_KEY:
        st.error("⚠️ Hệ thống chưa cấu hình API Key ở mục Secrets trên Streamlit Cloud!")
    elif not lesson_title:
        st.error("⚠️ Vui lòng điền tên bài học / chủ đề!")
    else:
        try:
            genai.configure(api_key=GOOGLE_API_KEY)
            model = genai.GenerativeModel("gemini-1.5-pro")

            prompt = f"""
            Bạn là một chuyên gia giáo dục phổ thông xuất sắc tại Việt Nam.
            Hãy soạn thảo tài liệu chuyên môn cho môn {subject}, {grade}, bộ sách {book}, thời lượng {duration} tiết.
            Bài học / Chủ đề: "{lesson_title}".

            Yêu cầu thực hiện chính xác các nội dung được chọn sau:
            """

            if "Đề kiểm tra định kỳ + Ma trận & Bản đặc tả (Chuẩn CV 7991)" in export_type:
                prompt += """
                1. ĐỀ KIỂM TRA ĐỊNH KỲ VÀ MA TRẬN, BẢN ĐẶC TẢ (TUÂN THỦ CHÍNH XÁC CÔNG VĂN 7991/BGDĐT-GDTrH):
                   A. Khung Ma trận & Bản đặc tả:
                      - Tỉ lệ mức độ nhận thức: 40% Nhận biết - 30% Thông hiểu - 30% Vận dụng.
                   B. Đề kiểm tra (Đúng định dạng 3 phần theo CV 7991):
                      - PHẦN I: Trắc nghiệm nhiều lựa chọn (4 phương án A, B, C, D - Khoảng 8-12 câu).
                      - PHẦN II: Trắc nghiệm Đúng - Sai (Gồm các câu lớn, mỗi câu có 4 ý a, b, c, d).
                      - PHẦN III: Câu hỏi trả lời ngắn (Điền kết quả ngắn gọn).
                   C. Đáp án & Hướng dẫn chấm chi tiết.
                """

            if "Kế hoạch bài dạy / Giáo án (Chuẩn CV 5512)" in export_type:
                prompt += """
                2. KẾ HOẠCH BÀI DẠY (TUÂN THỦ CÔNG VĂN 5512/BGDĐT-GDTrH):
                   - I. MỤC TIÊU (Kiến thức, Năng lực, Phẩm chất).
                   - II. THIẾT BỊ DẠY HỌC VÀ HỌC LIỆU.
                   - III. TIẾN TRÌNH DẠY HỌC (4 hoạt động: Khởi động, Hình thành kiến thức mới, Luyện tập, Vận dụng; mỗi hoạt động đủ 4 bước: Mục tiêu, Nội dung, Sản phẩm, Tổ chức thực hiện).
                """

            if "Kịch bản Slide bài giảng (PowerPoint)" in export_type:
                prompt += """
                3. KỊCH BẢN SLIDE BÀI GIẢNG (POWERPOINT):
                   - Chia rõ từng Slide: Tiêu đề, Nội dung hiển thị, Ghi chú giáo viên.
                """

            prompt += "\nTrình bày rõ ràng, chuẩn văn phong sư phạm bằng tiếng Việt."

            with st.spinner("🔄 AI đang phân tích chương trình và sinh nội dung chuẩn CV 7991..."):
                response = model.generate_content(prompt)
                output_text = response.text

            st.success("✅ Đã xuất bản thành công tài liệu từ AI!")
            st.markdown("---")
            st.markdown(output_text)

            # Nút Tải file Word (.docx)
            st.markdown("---")
            word_file = export_to_word(output_text)
            st.download_button(
                label="📥 Tải xuống File Word (.docx)",
                data=word_file,
                file_name=f"TaiLieu_CV7991_{subject}_{lesson_title}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                type="primary"
            )

        except Exception as e:
            st.error(f"❌ Có lỗi xảy ra khi gọi AI: {str(e)}")
